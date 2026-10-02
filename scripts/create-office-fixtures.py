"""Rebuild public deterministic fixtures; these libraries are fixture-only.

python -m pip install python-docx==1.2.0 python-pptx==1.0.2 openpyxl==3.1.5
python scripts/create-office-fixtures.py
Nothing here is a production dependency. No private documents are used.
"""
import io
import zipfile
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree as ET
from docx import Document
from docx.shared import Inches, Pt
from pptx import Presentation
from pptx.util import Inches as SlideInches, Pt as SlidePt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

root = Path(__file__).resolve().parents[1] / 'apps/api/src/tools/fixtures/office'
root.mkdir(exist_ok=True)
image = root.parent / 'english-scan.png'
fixed = datetime(2026, 1, 1)


def save(document, name):
    buffer = io.BytesIO()
    document.save(buffer)
    with zipfile.ZipFile(buffer) as source, zipfile.ZipFile(root / name, 'w', zipfile.ZIP_DEFLATED, compresslevel=6) as output:
        for part in sorted(source.namelist()):
            if 'printerSettings' in part:
                continue  # the bundled presentation template contains unused binary printer data
            data = source.read(part)
            if part.endswith('.rels') or part == '[Content_Types].xml':
                xml = ET.fromstring(data)
                # Keep each OPC root namespace as the default namespace. Some
                # LibreOffice import filters rely on the conventional package
                # serialization even though prefixed names are XML-equivalent.
                ET.register_namespace('', xml.tag.partition('}')[0].removeprefix('{'))
                for child in list(xml):
                    if 'printerSettings' in child.get('Type', '') or child.get('Extension') == 'bin':
                        xml.remove(child)
                data = ET.tostring(xml, encoding='utf-8', xml_declaration=True)
            info = zipfile.ZipInfo(part, (2026, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            output.writestr(info, data)


for with_image, name in [(False, 'paragraphs.docx'), (True, 'image-page-break.docx')]:
    doc = Document()
    doc.core_properties.created = doc.core_properties.modified = fixed
    style = doc.styles['Normal']
    style.font.name, style.font.size = 'Liberation Sans', Pt(12)
    doc.add_heading('Kagaz Writer Conversion', 0)
    doc.add_paragraph('First paragraph: Office conversion preserves meaningful content.')
    run = doc.add_paragraph().add_run('Bold formatted paragraph and blue image below.')
    run.bold = True
    if with_image:
        doc.add_picture(str(image), width=Inches(2))
        doc.add_page_break()
        doc.add_heading('Second Page', 1)
        doc.add_paragraph('Explicit page break produces another readable PDF page.')
    save(doc, name)

deck = Presentation()
deck.core_properties.created = deck.core_properties.modified = fixed
deck.slide_width, deck.slide_height = SlideInches(10), SlideInches(7.5)
for number in range(1, 4):
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    frame = slide.shapes.add_textbox(SlideInches(0.7), SlideInches(0.6), SlideInches(8), SlideInches(2)).text_frame
    frame.text = f'Kagaz Slide {number}'
    frame.paragraphs[0].runs[0].font.name = 'Liberation Sans'
    frame.paragraphs[0].runs[0].font.size = SlidePt(32)
    paragraph = frame.add_paragraph()
    paragraph.text = f'Meaningful presentation content on slide {number}.'
    paragraph.runs[0].font.size = SlidePt(18)
    if number == 2:
        slide.shapes.add_picture(str(image), SlideInches(1), SlideInches(3), height=SlideInches(3))
save(deck, 'slides.pptx')

for multiple, name in [(False, 'sheet.xlsx'), (True, 'sheets.xlsx')]:
    book = Workbook()
    book.properties.created = book.properties.modified = fixed
    book.remove(book.active)
    for title in (['Budget', 'Summary'] if multiple else ['Budget']):
        sheet = book.create_sheet(title)
        sheet.append([f'Kagaz {title}', 'Amount'])
        sheet.append(['Paper', 125])
        sheet.append(['Printing', 250])
        sheet.append(['Total', '=SUM(B2:B3)'])
        for cell in sheet[1]:
            cell.font = Font(name='Liberation Sans', size=14, bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='166B86')
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name='Liberation Sans', size=12)
                cell.alignment = Alignment(vertical='center')
        sheet.column_dimensions['A'].width = 32
        sheet.column_dimensions['B'].width = 18
        sheet.print_area = 'A1:B4'
        sheet.page_setup.orientation = 'portrait'
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = sheet.page_setup.fitToHeight = 1
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
    save(book, name)
